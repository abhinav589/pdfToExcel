import io
import os
import re
import sys
import shutil
import tempfile
import threading
import subprocess
import zipfile
from pathlib import Path

import streamlit as st
from openpyxl import load_workbook

APP_DIR = Path(__file__).resolve().parent
OCR_SCRIPT = APP_DIR / "279_pdf_to_excel.py"
WORK_DIR = APP_DIR / "work"
PROCESS_LOCK = threading.Lock()
MAX_FILES = 5
MAX_FILE_BYTES = 50 * 1024 * 1024
MAX_TOTAL_BYTES = 100 * 1024 * 1024

st.set_page_config(page_title="Voter PDF to Excel", page_icon="📄", layout="wide")
st.title("Voter list PDF to Excel")
st.write("Upload a Hindi/English voter-list PDF and download the extracted records as an Excel workbook.")
st.info(
    "This version is designed for voter-list PDFs with the same three-column, "
    "ten-card-row layout as the supplied example. Uploaded PDFs and OCR images "
    "are processed temporarily; do not upload files unless you are authorized "
    "to share them with this hosted service."
)

uploads = st.file_uploader(
    "Choose PDF file(s)",
    type=["pdf"],
    accept_multiple_files=True,
    help="Up to 5 files, 50 MB each, and 100 MB total per run.",
)

with st.expander("Processing options", expanded=True):
    first_page = st.number_input("First voter-list page", min_value=1, value=3, step=1)
    last_page = st.number_input(
        "Last voter-list page (0 means final PDF page)", min_value=0, value=0, step=1
    )
    expected_records = None
    if uploads and len(uploads) == 1:
        expected_value = st.number_input(
            "Expected record count (optional)",
            min_value=0,
            value=0,
            step=1,
            help="Leave at 0 to skip the count check. A mismatch produces a warning but still creates the workbook.",
        )
        if expected_value > 0:
            expected_records = int(expected_value)
    elif uploads and len(uploads) > 1:
        st.caption("For batch uploads, the expected-count check is omitted; each workbook reports its extracted count.")


def safe_stem(filename: str) -> str:
    stem = Path(filename).stem
    stem = re.sub(r"[^A-Za-z0-9._-]+", "_", stem).strip("._-")
    return stem[:80] or "voter_records"


def parse_record_count(output: str):
    match = re.search(r"\brecords\s+(\d+)\b", output)
    return int(match.group(1)) if match else None


def run_one(upload, index, total, first, last, expected):
    pdf_bytes = upload.getvalue()
    if not pdf_bytes.startswith(b"%PDF-"):
        raise ValueError(f"{upload.name}: file does not look like a PDF.")
    if len(pdf_bytes) > MAX_FILE_BYTES:
        raise ValueError(f"{upload.name}: exceeds the 50 MB per-file limit.")

    name = safe_stem(upload.name)
    with tempfile.TemporaryDirectory(prefix="voter_pdf_") as temp_name:
        temp_dir = Path(temp_name)
        input_pdf = temp_dir / (name + ".pdf")
        output_xlsx = temp_dir / (name + "_voter_records.xlsx")
        input_pdf.write_bytes(pdf_bytes)
        command = [sys.executable, str(OCR_SCRIPT), str(input_pdf), str(output_xlsx), "--first-page", str(first)]
        if last > 0:
            command.extend(["--last-page", str(last)])
        if expected is not None:
            command.extend(["--expected-records", str(expected)])

        # The existing CLI script uses one shared work/ directory beside itself.
        # Serialize runs and clear that temporary directory before and after each job.
        with PROCESS_LOCK:
            shutil.rmtree(WORK_DIR, ignore_errors=True)
            try:
                completed = subprocess.run(
                    command,
                    cwd=APP_DIR,
                    capture_output=True,
                    text=True,
                    timeout=1800,
                    check=False,
                )
                if completed.returncode != 0:
                    details = (completed.stderr or completed.stdout or "No diagnostic output.")[-4000:]
                    raise RuntimeError(f"OCR failed for {upload.name}.\n{details}")
                if not output_xlsx.is_file():
                    raise RuntimeError(f"The OCR process finished but did not create a workbook for {upload.name}.")
                workbook_bytes = output_xlsx.read_bytes()
                count = parse_record_count(completed.stdout)
                warnings = []
                if "WARNING:" in completed.stdout:
                    warnings = [line.strip() for line in completed.stdout.splitlines() if "WARNING:" in line]
            finally:
                shutil.rmtree(WORK_DIR, ignore_errors=True)

    return {
        "name": upload.name,
        "download_name": name + "_voter_records.xlsx",
        "bytes": workbook_bytes,
        "count": count,
        "warnings": warnings,
    }


if st.button("Extract and create Excel", type="primary", disabled=not uploads):
    st.session_state.pop("results", None)
    total_bytes = sum(len(f.getvalue()) for f in uploads)
    if len(uploads) > MAX_FILES:
        st.error(f"Choose no more than {MAX_FILES} PDFs per run.")
    elif total_bytes > MAX_TOTAL_BYTES:
        st.error("The combined upload size exceeds 100 MB.")
    else:
        results = []
        failures = []
        progress = st.progress(0, text="Starting OCR…")
        for i, upload in enumerate(uploads, start=1):
            progress.progress((i - 1) / len(uploads), text=f"Processing {upload.name} ({i}/{len(uploads)})…")
            try:
                with st.spinner(f"Running OCR for {upload.name}. Larger files can take several minutes."):
                    expected = expected_records if len(uploads) == 1 else None
                    result = run_one(upload, i, len(uploads), int(first_page), int(last_page), expected)
                results.append(result)
            except Exception as exc:
                failures.append(str(exc))
        progress.progress(1.0, text="Processing finished.")
        st.session_state["results"] = results
        st.session_state["failures"] = failures

if st.session_state.get("failures"):
    for error in st.session_state["failures"]:
        st.error(error)

results = st.session_state.get("results", [])
if results:
    st.success(f"Created {len(results)} workbook(s).")
    for i, result in enumerate(results):
        with st.expander(f"{result['name']} — {result['count'] if result['count'] is not None else 'record'} records", expanded=(len(results) == 1)):
            if result["count"] is not None:
                st.write(f"Extracted records: **{result['count']}**")
            for warning in result["warnings"]:
                st.warning(warning)
            # Show a small preview to let the user spot obvious OCR problems.
            preview_book = load_workbook(io.BytesIO(result["bytes"]), read_only=True, data_only=True)
            preview_sheet = preview_book["Voter Records"]
            preview_rows = list(preview_sheet.iter_rows(min_row=1, max_row=min(preview_sheet.max_row, 11), values_only=True))
            st.dataframe(preview_rows[1:], column_config={i: st.column_config.TextColumn(str(v)) for i, v in enumerate(preview_rows[0])}, hide_index=True)
            preview_book.close()
            st.download_button(
                "Download Excel",
                data=result["bytes"],
                file_name=result["download_name"],
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                key=f"download_{i}",
            )
    if len(results) > 1:
        zip_buffer = io.BytesIO()
        with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as archive:
            for result in results:
                archive.writestr(result["download_name"], result["bytes"])
        st.download_button(
            "Download all workbooks (ZIP)",
            data=zip_buffer.getvalue(),
            file_name="voter_records_workbooks.zip",
            mime="application/zip",
            key="download_all",
        )
    if st.button("Clear results from this session"):
        st.session_state.pop("results", None)
        st.session_state.pop("failures", None)
        st.rerun()
